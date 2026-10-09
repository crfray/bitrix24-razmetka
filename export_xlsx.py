# -*- coding: utf-8 -*-
"""Выгрузка спеки в xlsx: по листу на инструмент. Запуск: python export_xlsx.py tasks

Формул в файле нет намеренно: это снимок данных, а не модель. Проверить
формулы тут нечем (LibreOffice не установлен), а непроверенная формула
в отчёте опаснее готового числа.
"""
import collections
import datetime
import glob
import io
import json
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import check_anomalies
import space

SPACE = space.pick()
TODAY = datetime.date.today().isoformat()
ROOT = os.path.dirname(os.path.abspath(__file__))
EVENTS = os.path.join(SPACE["data"], "events")
OUT = os.path.join(os.path.expanduser("~"), "Desktop",
                   "Разметка 2.0 · %s — выгрузка %s.xlsx" % (SPACE["title"], TODAY))

FONT = "Arial"
STATE_LABEL = {
    "draft": "черновик", "approved": "согласование", "in_dev": "в разработке",
    "testing": "тестируется", "awaiting_release": "ждёт выкатки",
    "live": "работает сейчас", "broken": "сломано", "deprecated": "неактуально",
}
IMPL_LABEL = {
    "backlog": "Бэклог", "in_sprint": "В этом спринте", "rework": "На доработке",
    "needs_check": "Требует проверки", "needs_rework": "Требует доработки",
    "works": "Работает корректно",
}
PRIORITY_LABEL = {"high": "Высокий", "medium": "Средний", "low": "Низкий"}

# Та же шкала, что в интерфейсе: красный — далеко от готовности, зелёный — работает.
TONE_FILL = {
    1: "FAE5E2", 2: "FBEEDD", 3: "F8F0CF", 4: "F3F2CD", 5: "EAF3D7", 6: "E2F1E3",
    0: "ECEEF1",
}
STATE_TONE = {"draft": 1, "approved": 2, "in_dev": 3, "testing": 4,
              "awaiting_release": 5, "live": 6, "broken": 1, "deprecated": 0}
IMPL_TONE = {"backlog": 0, "needs_rework": 1, "in_sprint": 2, "rework": 3,
             "needs_check": 4, "works": 6}
PRIORITY_TONE = {"high": 1, "medium": 3, "low": 0}

COLUMNS = [
    ("важное", 8), ("Описание", 44), ("event", 26), ("приоритет", 12),
    ("состояние реализации", 20), ("состояние события", 18),
    ("кто сейчас занимается", 18), ("задача в Битрикс24", 26),
    ("фронт/бэк", 11), ("команда", 14), ("комментарий", 30),
    ("проверка разметки", 26), ("меры, в которых используется", 24),
    ("tool", 12), ("category", 18), ("type", 14),
    ("c_section", 22), ("c_sub_section", 22), ("c_element", 22), ("status", 12),
    ("p1", 18), ("p2", 20), ("p3", 18), ("p4", 16), ("p5", 22),
    ("код обновления", 16), ("testId", 22), ("выход", 12),
    ("скриншотов", 11), ("открытых вопросов", 12), ("вопросы", 46),
    ("источник", 18), ("id", 10),
]


def load_events():
    tools = {}
    for path in sorted(glob.glob(os.path.join(EVENTS, "*.json"))):
        data = json.load(io.open(path, encoding="utf-8"))
        tools[data["tool"]] = data["events"]
    return tools


def joined(values):
    return "\n".join(values or [])


def questions_text(event):
    lines = []
    for q in (event.get("questions") or []):
        mark = "[x]" if q.get("done") else "[ ]"
        line = mark + " " + q.get("text", "")
        if q.get("answer"):
            line += "  → " + q["answer"]
        lines.append(line)
    return "\n".join(lines)


def row_for(event):
    dev = event.get("dev") or {}
    params = event.get("params") or {}
    source = event.get("source") or {}
    return [
        "★" if event.get("starred") else "",
        event.get("description", ""),
        event.get("event", ""),
        PRIORITY_LABEL.get(event.get("priority"), ""),
        IMPL_LABEL.get(event.get("impl"), ""),
        STATE_LABEL.get(event.get("state"), event.get("state", "")),
        event.get("owner", ""),
        event.get("task_url", ""),
        event.get("platform", ""),
        event.get("team", ""),
        event.get("comment", ""),
        event.get("checked", ""),
        joined(event.get("measures")),
        event.get("tool", ""),
        event.get("category", ""),
        joined(event.get("type")) if isinstance(event.get("type"), list) else event.get("type", ""),
        joined(event.get("c_section")),
        joined(event.get("c_sub_section")),
        joined(event.get("c_element")),
        joined(event.get("status")),
        joined(params.get("p1")), joined(params.get("p2")), joined(params.get("p3")),
        joined(params.get("p4")), joined(params.get("p5")),
        dev.get("update_code", ""),
        joined(dev.get("test_ids")),
        dev.get("eta", ""),
        len(event.get("images") or []),
        len([q for q in (event.get("questions") or []) if not q.get("done")]),
        questions_text(event),
        "%s, стр %s" % (source.get("sheet", "?"), source.get("row", "?")),
        event.get("id", ""),
    ]


# Какие колонки несут пояснения к значениям — их вешаем примечанием к ячейке.
NOTE_COLUMNS = {
    "c_section": 17, "c_sub_section": 18, "c_element": 19, "status": 20,
    "p1": 21, "p2": 22, "p3": 23, "p4": 24, "p5": 25, "measures": 13,
}


def style_header(sheet):
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="E8EFFF")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = "A1:%s1" % get_column_letter(len(COLUMNS))


def write_tool_sheet(book, tool, events):
    sheet = book.create_sheet(tool[:31])
    style_header(sheet)
    events = sorted(events, key=lambda e: (not e.get("starred"), e.get("event", "")))
    for offset, event in enumerate(events):
        rownum = offset + 2
        for index, value in enumerate(row_for(event), start=1):
            cell = sheet.cell(row=rownum, column=index, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        # цвет статусов — тот же, что в интерфейсе
        for column, tone_map, key in ((4, PRIORITY_TONE, event.get("priority")),
                                      (5, IMPL_TONE, event.get("impl")),
                                      (6, STATE_TONE, event.get("state"))):
            tone = tone_map.get(key)
            if tone is not None:
                sheet.cell(row=rownum, column=column).fill = PatternFill(
                    "solid", fgColor=TONE_FILL[tone])
        # пояснения к значениям — примечанием, чтобы не плодить колонки
        for field, column in NOTE_COLUMNS.items():
            notes = (event.get("notes") or {}).get(field) or {}
            if not notes:
                continue
            text = "\n".join("%s — %s" % (k, v) for k, v in sorted(notes.items()))
            comment = Comment(text, "Разметка 2.0")
            comment.width = 340
            comment.height = 20 + 14 * min(10, len(notes))
            sheet.cell(row=rownum, column=column).comment = comment
    return sheet


def write_summary(book, tools):
    sheet = book.create_sheet("Сводка", 0)
    total = sum(len(v) for v in tools.values())
    rows = [
        ["Выгрузка спеки продуктовой разметки", ""],
        ["Дата выгрузки", TODAY],
        ["Источник", "%s/events/*.json инструмента «Разметка 2.0»" % os.path.basename(SPACE["data"])],
        ["", ""],
        ["Всего событий", total],
        ["Инструментов (листов)", len(tools)],
        ["", ""],
        ["Событий по инструментам", ""],
    ]
    for tool in sorted(tools, key=lambda t: -len(tools[t])):
        rows.append([tool, len(tools[tool])])
    rows.append(["", ""])

    flat = [e for events in tools.values() for e in events]
    rows.append(["Состояние реализации", ""])
    counter = collections.Counter(e.get("impl") or "" for e in flat)
    for key in ["backlog", "in_sprint", "rework", "needs_check", "needs_rework", "works"]:
        rows.append([IMPL_LABEL[key], counter.get(key, 0)])
    rows.append(["не указано", counter.get("", 0)])
    rows.append(["", ""])
    rows.append(["Открытых вопросов",
                 sum(len([q for q in (e.get("questions") or []) if not q.get("done")])
                     for e in flat)])
    rows.append(["Событий с открытыми вопросами",
                 sum(1 for e in flat
                     if any(not q.get("done") for q in (e.get("questions") or [])))])
    rows.append(["", ""])
    rows.append(["Значения перечислений разделены переносом строки внутри ячейки.", ""])
    rows.append(["Пояснения к значениям — в примечаниях к ячейкам (красный уголок).", ""])
    rows.append(["Формул в файле нет: это снимок данных, а не расчётная модель.", ""])

    for rownum, (left, right) in enumerate(rows, start=1):
        a = sheet.cell(row=rownum, column=1, value=left)
        b = sheet.cell(row=rownum, column=2, value=right)
        bold = left in ("Выгрузка спеки продуктовой разметки", "Событий по инструментам",
                        "Состояние реализации")
        a.font = Font(name=FONT, size=11, bold=bold)
        b.font = Font(name=FONT, size=11)
        a.alignment = Alignment(vertical="center")
    sheet.column_dimensions["A"].width = 62
    sheet.column_dimensions["B"].width = 18
    return sheet


def write_anomalies(book, events):
    found = check_anomalies.check(events)
    sheet = book.create_sheet("Аномалии")
    head = ["Что не так", "Событие", "Инструмент", "Подробности"]
    for index, title in enumerate(head, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor="FDF0E3")
    rownum = 2
    for kind in sorted(found, key=lambda k: -len(found[k])):
        for name, tool, detail in found[kind]:
            for index, value in enumerate([kind, name, tool, detail], start=1):
                cell = sheet.cell(row=rownum, column=index, value=value)
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            rownum += 1
    for letter, width in (("A", 44), ("B", 28), ("C", 14), ("D", 58)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    if rownum > 2:
        sheet.auto_filter.ref = "A1:D%d" % (rownum - 1)
    return rownum - 2


def main():
    tools = load_events()
    flat = [e for events in tools.values() for e in events]

    book = Workbook()
    book.remove(book.active)
    for tool in sorted(tools):
        write_tool_sheet(book, tool, tools[tool])
    write_summary(book, tools)
    anomalies = write_anomalies(book, flat)

    book.save(OUT)
    print("сохранено:", OUT)
    print("листов:", len(book.sheetnames), "→", ", ".join(book.sheetnames))
    print("событий:", len(flat))
    print("строк в «Аномалиях»:", anomalies)


if __name__ == "__main__":
    main()
